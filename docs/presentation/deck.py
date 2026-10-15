#!/usr/bin/env python3
"""One description of the 10 slides, written out twice: slides.html (for the PDF) and pvs-presentation.pptx (editable).

Every slide is a list of shapes placed on a 1280 x 720 px page (the same layout in both formats).
Run through docs/build.py, which also turns slides.html into the PDF.
"""
import html
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE.parent / "assets"
CACHE = HERE / "_img"
W, H = 1280, 720

PAGE, DARK, PANEL, PANEL2 = "#0f131c", "#0a0e17", "#1c1f29", "#262a34"
TEXT, SOFT, FAINT, LINE = "#dfe2ef", "#cbc3d7", "#958ea0", "#494454"
PURPLE, BLUE, GREEN, BLACK, WHITE = "#d0bcff", "#4cd7f6", "#4edea3", "#000000", "#ffffff"
HEAD, BODY, PIXEL, MONO = "Space Grotesk", "Inter", "Jersey 10", "DejaVu Sans Mono"
EVENT = "Afro-Asian Technology Competition 2026 · Games"


# ---------- shape helpers ----------
def rect(x, y, w, h, fill=None, border=None, bw=2, shadow=0):
    return dict(k="rect", x=x, y=y, w=w, h=h, fill=fill, border=border, bw=bw, shadow=shadow)


def text(x, y, w, h, s, size=20, color=TEXT, font=BODY, bold=False, align="left", spacing=0, lh=1.25, italic=False):
    return dict(k="text", x=x, y=y, w=w, h=h, s=s, size=size, color=color, font=font, bold=bold,
                align=align, spacing=spacing, lh=lh, italic=italic)


def image(path, x, y, w, h, border=None):
    """Crops the picture to the box (like CSS object-fit: cover) and keeps the result in _img/."""
    CACHE.mkdir(exist_ok=True)
    src = ASSETS / path
    out = CACHE / f"{src.stem}_{w}x{h}.png"
    im = Image.open(src).convert("RGBA")
    scale = max(w * 2 / im.width, h * 2 / im.height)          # 2x for sharp print
    im = im.resize((round(im.width * scale), round(im.height * scale)), Image.LANCZOS)
    left, top = (im.width - w * 2) // 2, (im.height - h * 2) // 2
    im.crop((left, top, left + w * 2, top + h * 2)).save(out)
    return dict(k="image", x=x, y=y, w=w, h=h, path=out, border=border)


def stripe(y, h=24):
    return [rect(i * h, y, h, h, fill=[PURPLE, BLUE, GREEN][i % 3]) for i in range(W // h + 1)]


def frame(n, title, bg=PAGE):
    """Title, footer and page number shared by the content slides."""
    return dict(bg=bg, shapes=[
        text(64, 40, 1152, 70, title, size=44, color=WHITE, font=HEAD, bold=True, lh=1.1),
        text(64, 682, 800, 24, f"Pixel Verse Studios · {EVENT}", size=13, color=FAINT),
        text(1152, 682, 64, 24, str(n), size=13, color=FAINT, align="right"),
    ])


def card(x, y, w, h, label, heading, body):
    return [rect(x, y, w, h, fill=PANEL, border=LINE, shadow=6),
            text(x + 32, y + 28, w - 64, 24, label.upper(), size=14, color=BLUE, font=HEAD, bold=True, spacing=3),
            text(x + 32, y + 64, w - 64, 90, heading, size=30, color=WHITE, font=HEAD, bold=True, lh=1.15),
            text(x + 32, y + 168, w - 64, h - 190, body, size=21, color=SOFT, lh=1.4)]


def bullets(x, y, w, items, gap=78):
    out = []
    for i, (strong, desc) in enumerate(items):
        yy = y + i * gap
        out += [rect(x, yy + 9, 12, 12, fill=GREEN),
                text(x + 28, yy, w - 28, 28, strong, size=21, color=WHITE, font=HEAD, bold=True),
                text(x + 28, yy + 30, w - 28, 40, desc, size=17, color=SOFT, lh=1.25)]
    return out


def pic_row(items, y, w, h, gap, x0=64, label_size=16):
    out = []
    for i, (path, cap) in enumerate(items):
        x = x0 + i * (w + gap)
        out += [image(path, x, y, w, h, border=LINE), text(x, y + h + 8, w, 26, cap, size=label_size, color=SOFT)]
    return out


# ---------- the slides ----------
def slides():
    out = []

    # 1. Title
    s = dict(bg=DARK, shapes=[
        text(64, 96, 1000, 28, EVENT.upper(), size=18, color=BLUE, font=HEAD, bold=True, spacing=4),
        text(64, 140, 760, 210, "Pixel Verse Studios", size=92, color=WHITE, font=HEAD, bold=True, lh=0.98),
        text(64, 372, 780, 60, "Real science, hidden in play.", size=44, color=PURPLE, font=PIXEL),
        text(64, 452, 640, 90, "Two 16-bit games for ages 5 to 99. Dodge ocean trash. Fly to Jupiter's icy moon. Every fact in them is real.",
             size=22, color=SOFT, lh=1.35),
        rect(880, 130, 300, 300, fill=PANEL, border=PURPLE, bw=4, shadow=10),
        image("brand/logo.png", 892, 142, 276, 276),
    ] + stripe(696))
    out.append(s)

    # 2. Problem and answer
    s = frame(2, "Facts stick when you play them.")
    s["shapes"] += card(64, 150, 544, 320, "The problem", "A fact in a lesson stays a sentence",
                        "\"Europa hides an ocean.\" \"Plastic hurts sea turtles.\" A child hears it, and the next day it is gone. Facts feel far away from play.")
    s["shapes"] += card(672, 150, 544, 320, "Our answer", "Make the fact a game moment",
                        "In Crab Quest you dodge the trash. In Space X you fly through Europa's ice. Each fact shows up where it makes sense in the story, and nothing in it is invented.")
    s["shapes"][-4] = rect(672, 150, 544, 320, fill=PANEL, border=PURPLE, shadow=6)
    s["shapes"].append(text(64, 520, 1152, 50, "Ages 5 to 99 · short sessions · runs offline on a school computer", size=22, color=PURPLE, font=HEAD, bold=True))
    out.append(s)

    # 3. Two games
    s = frame(3, "Two games, one studio.")
    for x, logo, name, sub, items in [
        (64, "brand/crab-quest-logo.png", "Crab Quest", "Six chapters to get home",
         [("6 chapters, 6 ways to play", "Platforming, dodging, collecting, surviving"),
          ("Real ocean threats", "Birds, trash and a predator fish"),
          ("Ends in a coral sanctuary", "Help a shrimp, then find your family")]),
        (672, "brand/space-x-logo.png", "Space X", "Journey to Jupiter",
         [("4 levels, one mission", "Launch pad, Earth orbit, asteroid belt, Europa"),
          ("A fact card after every level", "A story line and one science fact"),
          ("A planet fact on every loss", "The Sun and all eight planets, with pictures")])]:
        s["shapes"] += [rect(x, 140, 544, 500, fill=PANEL, border=LINE, shadow=6),
                        image(logo, x + 32, 168, 96, 96),
                        text(x + 150, 176, 360, 44, name, size=38, color=WHITE, font=HEAD, bold=True),
                        text(x + 150, 224, 360, 30, sub, size=20, color=PURPLE),
                        ] + bullets(x + 32, 300, 480, items, gap=108)
    out.append(s)

    # 4. Crab Quest
    s = frame(4, "Crab Quest: six chapters to get home")
    s["shapes"] += pic_row([("crab-quest/chapter-2.jpg", "Chapter 2 · Beach Run"),
                            ("crab-quest/chapter-4.jpg", "Chapter 4 · Shrimp Help"),
                            ("crab-quest/chapter-5.jpg", "Chapter 5 · Predator Chase")], 130, 370, 240, 21)
    for i, (n, name, line) in enumerate([("1", "Kitchen Escape", "Jump the counter"), ("2", "Beach Run", "Dodge birds, 30 s"),
                                         ("3", "Sea Floor Dodging", "Dodge trash, 10 s"), ("4", "Shrimp Help", "Collect 17 algae"),
                                         ("5", "Predator Chase", "Survive 45 s"), ("6", "Home Sweet Home", "Reach your family")]):
        x = 64 + i * 193
        s["shapes"] += [rect(x, 456, 181, 150, fill=PANEL2, border=LINE),
                        text(x + 14, 464, 60, 40, n, size=32, color=GREEN, font=PIXEL),
                        text(x + 14, 508, 155, 50, name, size=17, color=WHITE, font=HEAD, bold=True, lh=1.1),
                        text(x + 14, 560, 155, 40, line, size=14, color=SOFT)]
    out.append(s)

    # 5. Space X journey
    s = frame(5, "Space X: the Journey to Jupiter")
    s["shapes"] += pic_row([("space-x/launch_pad.png", "1 · Launch pad: collect 4 items"),
                            ("space-x/earth_orbit.png", "2 · Earth orbit: dodge 10 junk"),
                            ("space-x/belt.png", "3 · Asteroid belt: shoot 30"),
                            ("space-x/europa.png", "4 · Europa: break 12 ice chunks")], 130, 270, 203, 24, label_size=15)
    s["shapes"] += [rect(64, 408, 1152, 200, fill=PANEL, border=PURPLE, shadow=6),
                    text(96, 432, 1088, 26, "THE HERO'S MISSION", size=14, color=BLUE, font=HEAD, bold=True, spacing=3),
                    text(96, 468, 1088, 70, "Fly to Europa, Jupiter's icy moon, and find out if its hidden ocean could hold life.",
                         size=30, color=WHITE, font=HEAD, bold=True, lh=1.2),
                    text(96, 546, 1088, 50, "Lose all 3 hearts and the level restarts. Win level 4 and the Hero reaches Jupiter.", size=20, color=SOFT)]
    out.append(s)

    # 6. Learning moments
    s = frame(6, "Learning happens at the moment it matters")
    s["shapes"] += pic_row([("space-x/fact_card.png", "A fact card after every level"),
                            ("space-x/game_over.png", "A planet fact on every game over"),
                            ("space-x/win.png", "The win screen on Jupiter")], 130, 372, 279, 30)
    s["shapes"] += [text(64, 462, 1152, 40, "The Sun and all eight planets, each with a picture and a short fact.", size=24, color=WHITE, font=HEAD, bold=True),
                    text(64, 508, 1152, 80, "The facts live in plain Python lists, so adding a planet or a story line means adding one entry.", size=20, color=SOFT)]
    out.append(s)

    # 7. Architecture
    s = frame(7, "One simple idea: a mode switch")
    nodes = [("INPUT", "Keyboard · mouse", "arrows, Space, P, R, F and on-screen buttons"),
             ("PYGAME ZERO", "draw · update", "the only hooks in the main file"),
             ("ONE VARIABLE", "game mode", "says which screen is showing"),
             ("MODULES", "screens · levels", "each one owns its drawing and rules"),
             ("FILES", "images · music · data", "facts and targets in plain lists")]
    for i, (tag, name, small) in enumerate(nodes):
        x = 64 + i * 232
        hot = i == 2
        s["shapes"] += [rect(x, 150, 200, 170, fill="#2b2542" if hot else PANEL2, border=PURPLE if hot else LINE),
                        text(x + 12, 164, 176, 20, tag, size=12, color=BLUE, font=HEAD, bold=True, align="center", spacing=2),
                        text(x + 12, 194, 176, 60, name, size=23, color=WHITE, font=HEAD, bold=True, align="center", lh=1.1),
                        text(x + 14, 258, 172, 56, small, size=14, color=SOFT, align="center", lh=1.25)]
        if i < 4:
            s["shapes"].append(text(x + 200, 214, 32, 40, "›", size=40, color=PURPLE, font=HEAD, bold=True, align="center"))
    s["shapes"] += [rect(64, 372, 560, 230, fill=PANEL, border=LINE, shadow=6),
                    text(96, 392, 500, 28, "CRAB QUEST", size=14, color=BLUE, font=HEAD, bold=True, spacing=3),
                    text(96, 426, 500, 150, "START_SCREEN → LEVEL_SELECT → LEVEL1 … LEVEL6\n\nOne self-contained module per chapter, with its own sprites, platforms and reset().",
                         size=18, color=SOFT, lh=1.35),
                    rect(656, 372, 560, 230, fill=PANEL, border=LINE, shadow=6),
                    text(688, 392, 500, 28, "SPACE X", size=14, color=BLUE, font=HEAD, bold=True, spacing=3),
                    text(688, 426, 500, 150, "TITLE → PLAYING → FACT_CARD → GAME_OVER → WIN\n\nEach level decides when it is cleared. The main file only switches modes.",
                         size=18, color=SOFT, lh=1.35)]
    out.append(s)

    # 8. Tools and repo
    s = frame(8, "Built with free tools, in one repo")
    rows = [("Python 3.11+", "all game logic"), ("Pygame Zero 1.2.1", "sprites, keys, mouse, music"),
            ("Pygame 2.6.1", "fullscreen, sounds, images"), ("Jersey 10 font", "SIL Open Font License"),
            ("HTML + CSS", "the studio website"), ("PyInstaller · itch.io", "release path, in progress")]
    s["shapes"].append(rect(64, 140, 560, 480, fill=PANEL, border=LINE, shadow=6))
    for i, (a, b) in enumerate(rows):
        y = 164 + i * 74
        s["shapes"] += [text(96, y, 500, 30, a, size=24, color=WHITE, font=HEAD, bold=True),
                        text(96, y + 32, 500, 28, b, size=17, color=SOFT)]
    s["shapes"] += [rect(656, 140, 560, 480, fill=PAGE, border=PURPLE, shadow=6),
                    text(688, 164, 500, 28, "ONE GIT REPOSITORY", size=14, color=BLUE, font=HEAD, bold=True, spacing=3),
                    text(688, 204, 500, 380,
                         "PVS/\n├── Crab-game/            Crab Quest\n├── space-game/           Space X\n│   ├── screens/  levels/\n│   ├── parts/    data/\n│   └── images/ music/ fonts/\n├── pixel-verse-studios/  website\n└── docs/                 this deck",
                         size=19, color="#cfd3e4", font=MONO, lh=1.5)]
    out.append(s)

    # 9. Team
    s = frame(9, "Five people. Two games.")
    for i, (f, n) in enumerate([("mariam", "Mariam Alaa"), ("selim", "Selim Loai"), ("mohamed", "Mohamed Loai"),
                                ("medhat", "Medhat Ali"), ("salah", "Salah Mohamed")]):
        x = 64 + i * 234
        s["shapes"] += [rect(x, 160, 210, 270, fill=PANEL, border=PURPLE, shadow=6),
                        image(f"team/{f}.jpg", x + 10, 170, 190, 190),
                        text(x + 8, 374, 194, 40, n, size=22, color=WHITE, font=HEAD, bold=True, align="center")]
    s["shapes"].append(text(64, 480, 1152, 90, "Pixel Verse Studios: two complete educational games, with real facts and one clean architecture.",
                            size=28, color=PURPLE, font=HEAD, bold=True, lh=1.25))
    out.append(s)

    # 10. Closing
    s = dict(bg=DARK, shapes=[
        text(64, 80, 900, 100, "Thank you", size=88, color=WHITE, font=HEAD, bold=True),
        text(64, 200, 700, 40, "What's next", size=28, color=PURPLE, font=PIXEL),
    ] + bullets(64, 270, 700, [("Windows builds", "Package both games with PyInstaller"),
                               ("An itch.io page", "Both games, screenshots, controls and downloads"),
                               ("Test on a clean PC", "Every chapter and every level, music on and off")], gap=96) + [
        rect(880, 130, 300, 340, fill=WHITE),
        image("brand/qr-site.png", 892, 142, 276, 276),
        text(880, 424, 300, 30, "Scan to visit the studio", size=17, color="#14213D", align="center"),
        text(64, 600, 1000, 30, f"Pixel Verse Studios · {EVENT}", size=20, color=SOFT),
    ] + stripe(696))
    out.append(s)
    return out


# ---------- writers ----------
def css_color(c):
    return c


def to_html(deck):
    def one(sh):
        if sh["k"] == "rect":
            st = f"left:{sh['x']}px;top:{sh['y']}px;width:{sh['w']}px;height:{sh['h']}px;"
            if sh["fill"]:
                st += f"background:{sh['fill']};"
            if sh["border"]:
                st += f"border:{sh['bw']}px solid {sh['border']};box-sizing:border-box;"
            if sh["shadow"]:
                st += f"box-shadow:{sh['shadow']}px {sh['shadow']}px 0 #000;"
            return f'<div class="s" style="{st}"></div>'
        if sh["k"] == "image":
            st = f"left:{sh['x']}px;top:{sh['y']}px;width:{sh['w']}px;height:{sh['h']}px;"
            if sh["border"]:
                st += f"outline:2px solid {sh['border']};outline-offset:-2px;"
            return f'<img class="s" style="{st}" src="{Path(sh["path"]).relative_to(HERE).as_posix()}">'
        st = (f"left:{sh['x']}px;top:{sh['y']}px;width:{sh['w']}px;height:{sh['h']}px;font:{'italic ' if sh['italic'] else ''}"
              f"{700 if sh['bold'] else 400} {sh['size']}px/{sh['lh']} '{sh['font']}',sans-serif;color:{sh['color']};text-align:{sh['align']};"
              f"letter-spacing:{sh['spacing']}px;white-space:pre-wrap;")
        return f'<div class="s" style="{st}">{html.escape(sh["s"])}</div>'

    pages = "\n".join(f'<section class="slide" style="background:{sl["bg"]}">' + "".join(one(s) for s in sl["shapes"]) + "</section>"
                      for sl in deck)
    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Pixel Verse Studios: Presentation</title>
<link href="https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@500;700&family=Inter:wght@400;500;600&display=swap" rel="stylesheet">
<style>
@font-face {{ font-family:"Jersey 10"; src:url("../assets/brand/jersey10.ttf"); }}
@page {{ size:{W}px {H}px; margin:0; }}
* {{ margin:0; padding:0; }}
html {{ -webkit-print-color-adjust:exact; print-color-adjust:exact; }}
.slide {{ position:relative; width:{W}px; height:{H}px; overflow:hidden; page-break-after:always; break-after:page; }}
.s {{ position:absolute; display:block; }}
</style></head><body>{pages}</body></html>"""


def rgb(hexcolor):
    h = hexcolor.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def px(v):
    return Emu(round(v * 9525))


def to_pptx(deck, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    blank = prs.slide_layouts[6]
    for sl in deck:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(sl["bg"])
        for sh in sl["shapes"]:
            if sh["k"] == "rect":
                if sh["shadow"]:
                    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(sh["x"] + sh["shadow"]), px(sh["y"] + sh["shadow"]), px(sh["w"]), px(sh["h"]))
                    r.fill.solid(); r.fill.fore_color.rgb = rgb(BLACK); r.line.fill.background()
                r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(sh["x"]), px(sh["y"]), px(sh["w"]), px(sh["h"]))
                if sh["fill"]:
                    r.fill.solid(); r.fill.fore_color.rgb = rgb(sh["fill"])
                else:
                    r.fill.background()
                if sh["border"]:
                    r.line.color.rgb = rgb(sh["border"]); r.line.width = Pt(sh["bw"] * 0.75)
                else:
                    r.line.fill.background()
            elif sh["k"] == "image":
                p = slide.shapes.add_picture(str(sh["path"]), px(sh["x"]), px(sh["y"]), px(sh["w"]), px(sh["h"]))
                if sh["border"]:
                    p.line.color.rgb = rgb(sh["border"]); p.line.width = Pt(1.5)
            else:
                tb = slide.shapes.add_textbox(px(sh["x"]), px(sh["y"]), px(sh["w"]), px(sh["h"]))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, line in enumerate(sh["s"].split("\n")):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[sh["align"]]
                    para.line_spacing = sh["lh"]
                    run = para.add_run()
                    run.text = line
                    run.font.size = Pt(sh["size"] * 0.75)
                    run.font.bold = sh["bold"]
                    run.font.name = "Courier New" if sh["font"] == MONO else sh["font"]
                    run.font.color.rgb = rgb(sh["color"])
    prs.save(out)


if __name__ == "__main__":
    deck = slides()
    (HERE / "slides.html").write_text(to_html(deck), encoding="utf-8")
    to_pptx(deck, HERE / "pvs-presentation.pptx")
    print(len(deck), "slides written")
